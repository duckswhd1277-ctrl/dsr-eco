from openpyxl import load_workbook
import json

file1 = r"C:\Users\Admin\Desktop\정리파일(수출).xls"
file2 = r"C:\Users\Admin\Desktop\OT 열처리.xls"

try:
    wb1 = load_workbook(file1, data_only=True)
    sheet1 = wb1.active

    print("="*80)
    print("📄 정리파일(수출).xls")
    print("="*80)
    print(f"Sheet name: {sheet1.title}")

    # Get headers
    headers1 = []
    for cell in sheet1[1]:
        headers1.append(cell.value)

    print(f"컬럼: {headers1}")

    # Get data
    data1 = []
    for i, row in enumerate(sheet1.iter_rows(min_row=2, values_only=True), 1):
        data1.append(row)
        if i >= 5:
            break

    print("\n첫 5행 데이터:")
    for i, row in enumerate(data1, 1):
        print(f"행{i}: {row}")

    print("\n" + "="*80)
    print("📄 OT 열처리.xls")
    print("="*80)

    wb2 = load_workbook(file2, data_only=True)
    sheet2 = wb2.active

    print(f"Sheet name: {sheet2.title}")

    # Get headers
    headers2 = []
    for cell in sheet2[1]:
        headers2.append(cell.value)

    print(f"컬럼: {headers2}")

    # Get data
    data2 = []
    for i, row in enumerate(sheet2.iter_rows(min_row=2, values_only=True), 1):
        data2.append(row)
        if i >= 5:
            break

    print("\n첫 5행 데이터:")
    for i, row in enumerate(data2, 1):
        print(f"행{i}: {row}")

except Exception as e:
    print(f"Error: {e}")
    import traceback
    traceback.print_exc()
